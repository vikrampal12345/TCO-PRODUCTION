from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from backend.rag.loader import (
    DocumentLoaderError,
    load_file,
)


TEST_DIR = Path(__file__).resolve().parent
TEST_DATA_DIR = TEST_DIR / "data"
TEST_FILE = TEST_DATA_DIR / "pptx_loader_test.pptx"


def create_test_pptx() -> None:
    """
    Create a tiny PowerPoint file specifically for testing.
    """
    TEST_DATA_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    presentation = Presentation()

    slide = presentation.slides.add_slide(
        presentation.slide_layouts[5]
    )

    title = slide.shapes.title

    if title is not None:
        title.text = "Booth Algorithm"

    textbox = slide.shapes.add_textbox(
        Inches(1),
        Inches(2),
        Inches(8),
        Inches(2),
    )

    textbox.text_frame.text = (
        "Booth Algorithm is a multiplication "
        "algorithm for signed binary numbers."
    )

    slide_2 = presentation.slides.add_slide(
        presentation.slide_layouts[5]
    )

    title_2 = slide_2.shapes.title

    if title_2 is not None:
        title_2.text = "Booth Steps"

    textbox_2 = slide_2.shapes.add_textbox(
        Inches(1),
        Inches(2),
        Inches(8),
        Inches(2),
    )

    textbox_2.text_frame.text = (
        "Check Q0 and Q-1. "
        "Perform addition or subtraction. "
        "Then perform arithmetic shift right."
    )

    presentation.save(
        TEST_FILE
    )


def main() -> None:
    print()
    print("=" * 50)
    print("TCO PPTX LOADER TEST")
    print("=" * 50)

    create_test_pptx()

    print()
    print(
        f"Test PPTX: {TEST_FILE}"
    )

    assert TEST_FILE.exists()

    documents = load_file(
        str(TEST_FILE)
    )

    print()
    print(
        f"Slides loaded: {len(documents)}"
    )

    assert len(documents) == 2

    for document in documents:
        print()
        print(
            "Metadata:"
        )
        print(
            document["metadata"]
        )

        print()
        print(
            "Text:"
        )
        print(
            document["text"]
        )

        assert document["text"].strip()

        metadata = document[
            "metadata"
        ]

        assert metadata[
            "file_type"
        ] == "pptx"

        assert metadata[
            "file_name"
        ] == TEST_FILE.name

        assert metadata.get(
            "slide"
        ) in {1, 2}

    full_text = "\n".join(
        document["text"]
        for document in documents
    ).lower()

    assert "booth algorithm" in full_text
    assert "arithmetic shift right" in full_text

    print()
    print(
        "PPTX text extraction passed."
    )

    print()
    print("=" * 50)
    print("PPTX LOADER TEST PASSED")
    print("=" * 50)
    print()


if __name__ == "__main__":
    main()